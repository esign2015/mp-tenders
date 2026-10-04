import csv
import os
import hashlib
import hmac
import json
from datetime import datetime, timezone, timedelta
from pathlib import Path
from urllib import request as urllib_request, parse as urllib_parse
import subprocess
import tempfile
import threading
import time
from concurrent.futures import ThreadPoolExecutor

from flask import Flask, jsonify, request, Response
from flask_cors import CORS

from scraper import scrape_mp_tenders
import google_sheet_store as sheet_store

app = Flask(__name__)
CORS(app, expose_headers=["Server-Timing"])

@app.errorhandler(sheet_store.SheetStoreError)
def sheet_store_error(error):
    return jsonify({"ok":False,"message":str(error)}),error.status

ROOT = Path(__file__).resolve().parent.parent
CSV_FILE = ROOT / "all_tenders_org_detailed.csv"
FIELDS = [
    "Tender ID", "Published Date", "Closing Date", "Opening Date",
    "Title", "Reference Number", "Organisation", "Department",
    "Division", "Sub Division", "PAC Amount", "EMD Fee",
    "Tender Fee", "Processing Fee", "Total Fee", "Location", "Pincode",
    "Status", "URL",
]


_runtime_rows_cache={'at':0,'rows':None,'snapshot_at':None}
_runtime_rows_lock=threading.Lock()
def read_rows():
    if os.getenv('GOOGLE_SHEETS_WEBAPP_URL'):
        with _runtime_rows_lock:
            if _runtime_rows_cache['rows'] is not None and time.monotonic()-_runtime_rows_cache['at']<60:return _runtime_rows_cache['rows']
            try:
                import io,requests
                response=requests.get('https://raw.githubusercontent.com/esign2015/mp-tenders/tender-data/all_tenders_org_detailed.csv',timeout=20);response.raise_for_status()
                rows=list(csv.DictReader(io.StringIO(response.content.decode('utf-8-sig'))))
                _runtime_rows_cache.update(rows=rows,at=time.monotonic())
                return rows
            except Exception:
                if _runtime_rows_cache['rows'] is not None:return _runtime_rows_cache['rows']
    if not CSV_FILE.exists():return []
    with CSV_FILE.open('r',encoding='utf-8-sig',newline='') as f:return list(csv.DictReader(f))


@app.get("/")
def home():
    return jsonify({
        "status": "online",
        "message": "MP Tenders scraper API is running",
        "source": "https://www.mptenders.gov.in/nicgep/app",
    })


def clean(value):
    return str(value or "").strip()


# ================= USER REGISTRATION / ANALYTICS =================
# Telegram Login does not expose a user's phone number. We therefore ask the
# user for their name + mobile number after the first verified Telegram login.
# Data is stored in SQLite. Set USER_DB_PATH to a directory/file on a Render
# persistent disk for durable storage (for example /var/data/users.db).
USER_DB_PATH = Path(os.getenv("USER_DB_PATH", str(ROOT / "data" / "users.db")))
USER_DB_PATH.parent.mkdir(parents=True, exist_ok=True)
DATABASE_URL = clean(os.getenv("DATABASE_URL"))

class _DBRow(dict):
    def __getitem__(self, key):
        if isinstance(key, int):
            return list(self.values())[key]
        return super().__getitem__(key)

class _DBCursor:
    def __init__(self, cursor):
        self.cursor = cursor
    def fetchone(self):
        row = self.cursor.fetchone()
        if row is None:
            return None
        if isinstance(row, dict):
            return _DBRow(row)
        try:
            return _DBRow(dict(row))
        except Exception:
            return row
    def fetchall(self):
        rows = self.cursor.fetchall()
        out = []
        for row in rows:
            if isinstance(row, dict):
                out.append(_DBRow(row))
            else:
                try:
                    out.append(_DBRow(dict(row)))
                except Exception:
                    out.append(row)
        return out

class _DBConnection:
    def __init__(self, conn, postgres=False):
        self.conn = conn
        self.postgres = postgres
    def execute(self, sql, params=()):
        if self.postgres:
            sql = sql.replace("?", "%s")
        return _DBCursor(self.conn.execute(sql, params))
    def executescript(self, script):
        if self.postgres:
            for statement in script.split(";"):
                statement = statement.strip()
                if statement:
                    self.conn.execute(statement)
        else:
            self.conn.executescript(script)
    def commit(self):
        self.conn.commit()
    def close(self):
        self.conn.close()

IST = timezone(timedelta(hours=5, minutes=30))

def user_db():
    if DATABASE_URL:
        import psycopg
        from psycopg.rows import dict_row
        raw = psycopg.connect(DATABASE_URL, row_factory=dict_row, connect_timeout=15)
        conn = _DBConnection(raw, postgres=True)
        conn.executescript("""
            CREATE TABLE IF NOT EXISTS users (
                telegram_id BIGINT PRIMARY KEY,
                first_name TEXT NOT NULL DEFAULT '',
                last_name TEXT NOT NULL DEFAULT '',
                username TEXT NOT NULL DEFAULT '',
                name TEXT NOT NULL DEFAULT '',
                mobile TEXT NOT NULL DEFAULT '',
                mobile_verified INTEGER NOT NULL DEFAULT 0,
                signup_at TEXT NOT NULL,
                last_login_at TEXT NOT NULL,
                login_count INTEGER NOT NULL DEFAULT 0,
                email TEXT NOT NULL DEFAULT '',
                district TEXT NOT NULL DEFAULT '',
                state TEXT NOT NULL DEFAULT ''
            );
            CREATE TABLE IF NOT EXISTS login_events (
                id BIGSERIAL PRIMARY KEY,
                telegram_id BIGINT NOT NULL,
                login_at TEXT NOT NULL
            );
            CREATE INDEX IF NOT EXISTS idx_login_events_time ON login_events(login_at);
        """)
        conn.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS email TEXT NOT NULL DEFAULT ''")
        conn.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS district TEXT NOT NULL DEFAULT ''")
        conn.execute("ALTER TABLE users ADD COLUMN IF NOT EXISTS state TEXT NOT NULL DEFAULT ''")
        conn.commit()
        return conn

    import sqlite3
    raw = sqlite3.connect(str(USER_DB_PATH), timeout=30)
    raw.row_factory = sqlite3.Row
    raw.execute("PRAGMA journal_mode=WAL")
    raw.execute("PRAGMA busy_timeout=30000")
    conn = _DBConnection(raw, postgres=False)
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS users (
            telegram_id INTEGER PRIMARY KEY,
            first_name TEXT NOT NULL DEFAULT '',
            last_name TEXT NOT NULL DEFAULT '',
            username TEXT NOT NULL DEFAULT '',
            name TEXT NOT NULL DEFAULT '',
            mobile TEXT NOT NULL DEFAULT '',
            mobile_verified INTEGER NOT NULL DEFAULT 0,
            signup_at TEXT NOT NULL,
            last_login_at TEXT NOT NULL,
            login_count INTEGER NOT NULL DEFAULT 0,
            email TEXT NOT NULL DEFAULT '',
            district TEXT NOT NULL DEFAULT ''
        );
        CREATE TABLE IF NOT EXISTS login_events (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            telegram_id INTEGER NOT NULL,
            login_at TEXT NOT NULL
        );
        CREATE INDEX IF NOT EXISTS idx_login_events_time
            ON login_events(login_at);
    """)
    existing = {row[1] for row in raw.execute("PRAGMA table_info(users)").fetchall()}
    if "email" not in existing:
        raw.execute("ALTER TABLE users ADD COLUMN email TEXT NOT NULL DEFAULT ''")
    if "district" not in existing:
        raw.execute("ALTER TABLE users ADD COLUMN district TEXT NOT NULL DEFAULT ''")
    if "state" not in existing:
        raw.execute("ALTER TABLE users ADD COLUMN state TEXT NOT NULL DEFAULT ''")
    raw.commit()
    return conn

def user_db_backend():
    return "postgres" if DATABASE_URL else "sqlite"

def now_ist():
    return datetime.now(IST)

def normalise_mobile(value):
    digits = "".join(ch for ch in clean(value) if ch.isdigit())
    if digits.startswith("91") and len(digits) == 12:
        digits = digits[2:]
    if len(digits) != 10 or digits[0] not in "6789":
        return ""
    return "+91" + digits

def admin_telegram_ids():
    raw = clean(os.getenv("ADMIN_TELEGRAM_IDS", ""))
    ids = set()
    for item in raw.split(","):
        item = clean(item)
        if item.isdigit():
            ids.add(int(item))
    return ids

def is_user_admin(user_id):
    if int(user_id or 0) in admin_telegram_ids():
        return True
    # The channel owner (creator) may access the user report without another
    # secret being placed in the frontend.
    try:
        channel = clean(os.getenv("TELEGRAM_CHANNEL", "@mptendersalert"))
        member = telegram_api("getChatMember", {"chat_id": channel, "user_id": int(user_id)})
        return bool(member.get("ok")) and member.get("result", {}).get("status") == "creator"
    except Exception:
        return False

def touch_user_login(user_id, telegram_payload=None):
    telegram_payload = telegram_payload or {}
    uid = int(user_id)
    now = now_ist().isoformat()
    first_name = clean(telegram_payload.get("first_name"))
    last_name = clean(telegram_payload.get("last_name"))
    username = clean(telegram_payload.get("username"))
    display_name = " ".join(x for x in (first_name, last_name) if x).strip()

    conn = user_db()
    try:
        row = conn.execute("SELECT * FROM users WHERE telegram_id=?", (uid,)).fetchone()
        if row is None:
            conn.execute("""
                INSERT INTO users
                (telegram_id, first_name, last_name, username, name, mobile,
                 mobile_verified, signup_at, last_login_at, login_count)
                VALUES (?, ?, ?, ?, ?, '', 0, ?, ?, 1)
            """, (uid, first_name, last_name, username, display_name, now, now))
        else:
            conn.execute("""
                UPDATE users
                SET first_name=COALESCE(NULLIF(?, ''), first_name),
                    last_name=COALESCE(NULLIF(?, ''), last_name),
                    username=COALESCE(NULLIF(?, ''), username),
                    last_login_at=?,
                    login_count=login_count+1
                WHERE telegram_id=?
            """, (first_name, last_name, username, now, uid))
        conn.execute(
            "INSERT INTO login_events (telegram_id, login_at) VALUES (?, ?)",
            (uid, now),
        )
        conn.commit()
        row = conn.execute("SELECT * FROM users WHERE telegram_id=?", (uid,)).fetchone()
        return dict(row)
    finally:
        conn.close()

def get_user(user_id):
    conn = user_db()
    try:
        row = conn.execute("SELECT * FROM users WHERE telegram_id=?", (int(user_id),)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()

def update_user_profile(user_id, name, mobile, email="", state="", district=""):
    import re
    name = clean(name)
    mobile = normalise_mobile(mobile)
    email = clean(email).lower()
    state = clean(state)
    district = clean(district)
    if len(name) < 2:
        raise ValueError("कृपया अपना पूरा नाम दर्ज करें।")
    if not mobile:
        raise ValueError("कृपया सही नंबर डालिए।")
    if not re.fullmatch(r"[^\s@]+@[^\s@]+\.[^\s@]{2,}", email):
        raise ValueError("कृपया सही ईमेल आईडी डालिए।")
    allowed_domains = {"gmail.com", "yahoo.com", "yahoo.co.in", "rediffmail.com"}
    local, domain = email.rsplit("@", 1) if "@" in email else ("", "")
    if not re.fullmatch(r"[A-Za-z0-9._%+-]+", local) or not re.fullmatch(r"[A-Za-z0-9.-]+\.[A-Za-z]{2,}", domain):
        raise ValueError("कृपया सही ईमेल आईडी डालिए।")
    if not state:
        raise ValueError("कृपया राज्य चुनें।")
    if len(district) < 2:
        raise ValueError("कृपया जिला चुनें।")
    conn = user_db()
    try:
        conn.execute("""
            UPDATE users
            SET name=?, mobile=?, email=?, state=?, district=?, mobile_verified=0
            WHERE telegram_id=?
        """, (name, mobile, email, state, district, int(user_id)))
        conn.commit()
        row = conn.execute("SELECT * FROM users WHERE telegram_id=?", (int(user_id),)).fetchone()
        return dict(row) if row else None
    finally:
        conn.close()

def user_stats():
    conn = user_db()
    try:
        total = conn.execute("SELECT COUNT(*) FROM users").fetchone()[0]
        registered_mobile = conn.execute(
            "SELECT COUNT(*) FROM users WHERE mobile<>''"
        ).fetchone()[0]
        today = now_ist().date().isoformat()
        today_new = conn.execute(
            "SELECT COUNT(*) FROM users WHERE substr(signup_at,1,10)=?", (today,)
        ).fetchone()[0]
        today_users = conn.execute(
            "SELECT COUNT(DISTINCT telegram_id) FROM login_events WHERE substr(login_at,1,10)=?",
            (today,),
        ).fetchone()[0]
        today_logins = conn.execute(
            "SELECT COUNT(*) FROM login_events WHERE substr(login_at,1,10)=?",
            (today,),
        ).fetchone()[0]
        daily_rows = conn.execute("""
            SELECT substr(signup_at,1,10) AS day, COUNT(*) AS new_users
            FROM users
            GROUP BY substr(signup_at,1,10)
            ORDER BY day DESC
            LIMIT 90
        """).fetchall()
        return {
            "total_users": total,
            "registered_mobile": registered_mobile,
            "today_new_users": today_new,
            "today_unique_users": today_users,
            "today_logins": today_logins,
            "daily_signups": [dict(r) for r in daily_rows],
        }
    finally:
        conn.close()

def visitor_verification_flags(visitors):
    # Password-account verification is canonical; legacy visitor rows may not
    # have this field. Query once per mobile only during the admin export.
    mobiles = {row['mobile'] for row in visitors}
    if not sheet_store.enabled() and not DATABASE_URL and os.getenv('ALLOW_EPHEMERAL_ACCOUNTS','0') != '1':
        return {mobile:any(row['mobile']==mobile and str(dict(row).get('mobile_verified','')).lower() in {'true','yes','1'} for row in visitors) for mobile in mobiles}
    def read(mobile):
        record = account_service.get(mobile=mobile)
        legacy_verified = any(row['mobile'] == mobile and str(dict(row).get('mobile_verified', '')).lower() in {'true','yes','1'} for row in visitors)
        return mobile, legacy_verified or bool(record and record.get('mobile_verified') is True)
    with ThreadPoolExecutor(max_workers=4) as pool:
        return dict(pool.map(read, mobiles))

def build_user_excel():
    from io import BytesIO
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    conn = visitor_db()
    try:
        users = conn.execute("""
            SELECT telegram_id, name, mobile, email, state, district, username, first_name, last_name,
                   signup_at, last_login_at, login_count, mobile_verified
            FROM users
            ORDER BY signup_at DESC
        """).fetchall()
        visitors = (sheet_store.call('list_visitors')['visitors'] if sheet_store.enabled()
                    else conn.execute("SELECT * FROM visitor_registrations ORDER BY signup_at DESC").fetchall())
        daily = conn.execute("""
            SELECT substr(signup_at,1,10) AS day,
                   COUNT(*) AS new_users,
                   (SELECT COUNT(DISTINCT le.telegram_id)
                    FROM login_events le
                    WHERE substr(le.login_at,1,10)=substr(u.signup_at,1,10)) AS unique_logins,
                   (SELECT COUNT(*)
                    FROM login_events le2
                    WHERE substr(le2.login_at,1,10)=substr(u.signup_at,1,10)) AS total_logins
            FROM users u
            GROUP BY substr(signup_at,1,10)
            ORDER BY day DESC
        """).fetchall()
    finally:
        conn.close()

    verified_visitors = visitor_verification_flags(visitors)
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"
    headers = [
        "S.No.", "Name", "Mobile", "Email", "State", "District", "Telegram Username", "Telegram ID",
        "First Signup (IST)", "Last Login (IST)", "Login Count",
        "Mobile Verified"
    ]
    ws.append(headers)
    for i, row in enumerate(users, 1):
        ws.append([
            i, row["name"], row["mobile"], row["email"], row["state"], row["district"], ("@" + row["username"]) if row["username"] else "",
            row["telegram_id"], row["signup_at"], row["last_login_at"],
            row["login_count"], "Yes" if row["mobile_verified"] else "No"
        ])
    for cell in ws[1]:
        cell.font = cell.font.copy(bold=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    ds = wb.create_sheet("Daily Stats")
    ds.append(["Date (IST)", "New Users", "Unique Users", "Total Logins"])
    for row in daily:
        ds.append([row["day"], row["new_users"], row["unique_logins"], row["total_logins"]])
    for cell in ds[1]:
        cell.font = cell.font.copy(bold=True)
    ds.freeze_panes = "A2"

    vs = wb.create_sheet("Visitor Registrations")
    vs.append(["S.No.", "Name", "Mobile", "District", "First Saved (IST)", "Last Visit (IST)", "Visit Count", "Mobile Verified"])
    for i, row in enumerate(visitors, 1):
        vs.append([i,row["name"],row["mobile"],row["district"],row["signup_at"],row["last_visit_at"],row["visit_count"],"Yes" if verified_visitors.get(row["mobile"]) else "No"])
    for row in vs.iter_rows(min_row=2):
        for cell in row[1:4]:
            cell.data_type = "s"
    for cell in vs[1]:
        cell.font = cell.font.copy(bold=True)
    vs.freeze_panes = "A2"
    vs.auto_filter.ref = vs.dimensions
    for sheet in (ws, ds, vs):
        for col in range(1, sheet.max_column + 1):
            max_len = max(
                len(str(sheet.cell(row=r, column=col).value or ""))
                for r in range(1, min(sheet.max_row, 200) + 1)
            )
            sheet.column_dimensions[get_column_letter(col)].width = min(max(max_len + 2, 12), 34)

    output = BytesIO()
    wb.save(output)
    output.seek(0)
    return output.read()

# Temporary public access: collect a profile without Telegram authentication.
def visitor_db():
    conn = user_db()
    conn.executescript("""
        CREATE TABLE IF NOT EXISTS visitor_registrations (
            visitor_id TEXT PRIMARY KEY,
            name TEXT NOT NULL,
            mobile TEXT NOT NULL,
            district TEXT NOT NULL,
            signup_at TEXT NOT NULL,
            last_visit_at TEXT NOT NULL,
            visit_count INTEGER NOT NULL DEFAULT 1
        );
        CREATE TABLE IF NOT EXISTS visitor_events (
            event_id TEXT PRIMARY KEY,
            visitor_id TEXT NOT NULL,
            visited_at TEXT NOT NULL
        );
        CREATE TABLE IF NOT EXISTS visitor_affidavit_profiles (
            visitor_id TEXT PRIMARY KEY,
            profile_json TEXT NOT NULL,
            updated_at TEXT NOT NULL
        );
    """)
    conn.commit()
    return conn


@app.get("/api/admin/visitor-registrations")
def admin_visitor_registrations():
    email, _ = require_admin()
    if not email:
        return jsonify({"ok":False,"message":"Admin access required."}),401
    if sheet_store.enabled():
        response=jsonify({"ok":True,"visitors":sheet_store.call('list_visitors')['visitors']})
        response.headers['Cache-Control']='no-store'
        return response
    conn = visitor_db()
    try:
        rows = conn.execute("SELECT * FROM visitor_registrations ORDER BY signup_at DESC").fetchall()
        response = jsonify({"ok":True,"visitors":[dict(row) for row in rows]})
        response.headers["Cache-Control"] = "no-store"
        return response
    finally:
        conn.close()


@app.post('/api/visitors/affidavit')
def visitor_affidavit():
    payload=request.get_json(silent=True) or {}
    token=payload.get('session_token')
    account=account_service.authenticate(token)
    visitor_id=account['user_id']
    if not visitor_id:
        return jsonify({'ok':False,'message':'Saved profile session required.'}),401
    profile=payload.get('profile')
    if profile is not None:
        keys=('bidderName', 'parentRelation', 'parentName', 'address', 'houseNumber', 'roadStreet', 'locality', 'landmark', 'pincode', 'district', 'tehsil', 'firmName', 'status', 'place', 'relative', 'relativeName', 'relativePost', 'relativePosting', 'email', 'registrationNumber', 'registrationClass', 'registrationDate', 'registrationValidTill', 'pan', 'gst', 'telephone', 'fax', 'representativeName', 'representativeDesignation', 'representativeAddress', 'representativeTelephone', 'representativeFax', 'representativeMobile', 'representativeEmail', 'organisationType')
        if not isinstance(profile,dict):
            return jsonify({'ok':False,'message':'Invalid affidavit profile.'}),400
        saved_profile=(account or {}).get('affidavit_profile') or {}
        if not saved_profile and not sheet_store.enabled():
            existing=visitor_db()
            try:
                row=existing.execute('SELECT profile_json FROM visitor_affidavit_profiles WHERE visitor_id=?',(visitor_id,)).fetchone()
                saved_profile=json.loads(row['profile_json']) if row else {}
            finally:
                existing.close()
        address_keys=('houseNumber','roadStreet','locality','landmark')
        structured_address=any(key in profile for key in address_keys)
        old_client_address=profile.get('address')
        profile={key:clean(profile.get(key,saved_profile.get(key))) for key in keys}
        if structured_address:
            profile['address']=', '.join(profile[key] for key in address_keys if profile[key])
            if not profile['address']:
                return jsonify({'ok':False,'message':'Firm Address भरें.'}),400
        elif old_client_address is not None and clean(old_client_address)!=clean(saved_profile.get('address')):
            for key in address_keys:profile[key]=''
            profile['roadStreet']=profile['address']
        profile['parentRelation']=profile['parentRelation'] or 'S/o'
        import re
        profile['gst']=profile['gst'].upper();profile['pan']=profile['pan'].upper()
        if profile['gst'] and not re.fullmatch(r'[0-9]{2}[A-Z]{5}[0-9]{4}[A-Z][A-Z0-9]{3}',profile['gst']):
            return jsonify({'ok':False,'message':'GST में सही क्रम के 15 characters भरें.'}),400
        if profile['gst'] and not profile['pan']:profile['pan']=profile['gst'][2:12]
        if profile['pan'] and not re.fullmatch(r'[A-Z]{5}[0-9]{4}[A-Z]',profile['pan']):
            return jsonify({'ok':False,'message':'PAN में 5 अक्षर, 4 अंक और 1 अक्षर (कुल 10) भरें.'}),400
        if profile['pincode'] and not re.fullmatch(r'[1-9][0-9]{5}',profile['pincode']):
            return jsonify({'ok':False,'message':'Pincode में सही 6 अंक भरें.'}),400
        if not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',profile['email']):
            return jsonify({'ok':False,'message':'Letterhead के लिए सही email अनिवार्य है.'}),400
        if profile['representativeEmail'] and not re.fullmatch(r'[^\s@]+@[^\s@]+\.[^\s@]+',profile['representativeEmail']):
            return jsonify({'ok':False,'message':'Authorized representative का email सही भरें.'}),400
        for key in ('registrationDate','registrationValidTill'):
            if profile[key]:
                try:
                    from datetime import date
                    date.fromisoformat(profile[key])
                except ValueError:
                    return jsonify({'ok':False,'message':'Registration date सही भरें.'}),400

        if (any(len(value)>240 for value in profile.values())
            or not all(profile[key] for key in ('bidderName','firmName','status','place'))
            or profile['parentRelation'] not in ('S/o','D/o','W/o')
            or profile['relative'] not in ('yes','no')
            or (profile['relative']=='yes' and not all(profile[key] for key in ('relativeName','relativePost','relativePosting')))):
            return jsonify({'ok':False,'message':'Required affidavit basic details are missing or invalid.'}),400
    remote=sheet_store.enabled()
    if profile is not None:
        from bidder_tools import keep_backup
        for attempt in range(3):
            if attempt:account=account_service.authenticate(token)
            revision=account['revision'];keep_backup(account)
            account['affidavit_profile']=profile;account['affidavit_profile_updated_at']=now_ist().isoformat()
            if account_service.update(account,revision):break
        else:return jsonify({'ok':False,'message':'Profile changed during save. Please retry.'}),409
        if remote:
            from account_access import queue_affidavit_mirror
            mirrored=queue_affidavit_mirror(visitor_id,profile)
        else:
            mirrored=False
            conn=visitor_db()
            try:
                conn.execute('INSERT INTO visitor_affidavit_profiles (visitor_id,profile_json,updated_at) VALUES (?,?,?) ON CONFLICT(visitor_id) DO UPDATE SET profile_json=excluded.profile_json,updated_at=excluded.updated_at',(visitor_id,json.dumps(profile,ensure_ascii=False),now_ist().isoformat()));conn.commit()
            finally:conn.close()
        response=jsonify({'ok':True,'profile':profile,'canonical':True,'storage':'google_sheets' if remote else user_db_backend(),'affidavit_sheet_sync_pending':mirrored})
    elif isinstance(account.get('affidavit_profile'),dict):
        response=jsonify({'ok':True,'profile':account['affidavit_profile'],'canonical':True,'storage':'google_sheets' if remote else user_db_backend()})
    elif remote:
        result=sheet_store.call('read_affidavit',visitor_id=visitor_id)
        response=jsonify({'ok':True,'profile':result['profile'],'storage':'google_sheets'})
    else:
        conn=visitor_db()
        try:row=conn.execute('SELECT profile_json FROM visitor_affidavit_profiles WHERE visitor_id=?',(visitor_id,)).fetchone()
        finally:conn.close()
        response=jsonify({'ok':True,'profile':json.loads(row['profile_json']) if row else {},'storage':user_db_backend()})
    response.headers['Cache-Control']='no-store';return response


@app.post('/api/users/export')
def users_export():
    email,_=require_admin()
    if not email:return jsonify({'ok':False,'message':'Admin access required.'}),403
    return Response(build_user_excel(),mimetype='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="MP_Tender_Users_Report.xlsx"','Cache-Control':'no-store'})


@app.get('/api/admin/profile-storage')
def admin_profile_storage():
    email,_=require_admin()
    if not email:
        return jsonify({'ok':False,'message':'Admin access required.'}),401
    result=sheet_store.call('status') if sheet_store.enabled() else {}
    response=jsonify({'ok':True,'storage':'google_sheets' if sheet_store.enabled() else user_db_backend(),**result})
    response.headers['Cache-Control']='no-store'
    return response


# ================= GOOGLE ADMIN CONTROL PANEL =================
ADMIN_EMAIL_ALLOWLIST = {
    x.strip().lower()
    for x in clean(os.getenv(
        "ADMIN_GOOGLE_EMAILS",
        "imriteshdhoot@gmail.com,shikhadhoot@gmail.com",
    )).split(",")
    if x.strip()
}
ADMIN_SESSION_TTL = int(os.getenv("ADMIN_SESSION_TTL", "28800"))  # 8 hours

def admin_session_secret():
    return clean(os.getenv("ADMIN_SESSION_SECRET")) or clean(os.getenv("TELEGRAM_SESSION_SECRET")) or clean(os.getenv("TELEGRAM_BOT_TOKEN"))

def make_admin_session(email):
    import base64, time
    payload = {"email": email.lower(), "exp": int(time.time()) + ADMIN_SESSION_TTL}
    raw = json.dumps(payload, separators=(",", ":"), sort_keys=True).encode("utf-8")
    body = base64.urlsafe_b64encode(raw).decode("ascii").rstrip("=")
    sig = hmac.new(admin_session_secret().encode("utf-8"), body.encode("ascii"), hashlib.sha256).hexdigest()
    return body + "." + sig

def read_admin_session(token):
    import base64, time
    if not token or "." not in token:
        return None
    body, received_sig = token.rsplit(".", 1)
    expected = hmac.new(admin_session_secret().encode("utf-8"), body.encode("ascii"), hashlib.sha256).hexdigest()
    if not hmac.compare_digest(expected, received_sig):
        return None
    try:
        padded = body + "=" * (-len(body) % 4)
        payload = json.loads(base64.urlsafe_b64decode(padded.encode("ascii")).decode("utf-8"))
        if int(payload.get("exp", 0)) < int(time.time()):
            return None
        email = clean(payload.get("email")).lower()
        return email if email in ADMIN_EMAIL_ALLOWLIST else None
    except Exception:
        return None

def require_admin():
    token = clean(request.headers.get("Authorization", "")).removeprefix("Bearer ").strip()
    email = read_admin_session(token)
    if not email:
        return None, (jsonify({"ok": False, "message": "Admin login required."}), 401)
    return email, None

@app.get("/api/admin/config")
def admin_config():
    return jsonify({
        "ok": True,
        "google_client_id": clean(os.getenv("GOOGLE_CLIENT_ID")),
        "allowed_domains": ["google.com"],
    })

@app.post("/api/admin/google")
def admin_google_login():
    payload = request.get_json(silent=True) or {}
    credential = clean(payload.get("credential"))
    if not credential:
        return jsonify({"ok": False, "message": "Google authentication token missing."}), 400
    client_id = clean(os.getenv("GOOGLE_CLIENT_ID"))
    if not client_id:
        return jsonify({"ok": False, "message": "Google Admin login is not configured on the server yet."}), 503

    # Verify the Google ID token server-side. The tokeninfo endpoint validates
    # the signature and exposes aud/email/email_verified for our checks.
    try:
        url = "https://oauth2.googleapis.com/tokeninfo?" + urllib_parse.urlencode({"id_token": credential})
        with urllib_request.urlopen(url, timeout=20) as response:
            info = json.loads(response.read().decode("utf-8"))
    except Exception:
        return jsonify({"ok": False, "message": "Google login verification failed."}), 401

    email = clean(info.get("email")).lower()
    if clean(info.get("aud")) != client_id:
        return jsonify({"ok": False, "message": "Google client verification failed."}), 401
    if str(info.get("email_verified", "")).lower() != "true":
        return jsonify({"ok": False, "message": "Google email is not verified."}), 403
    if email not in ADMIN_EMAIL_ALLOWLIST:
        return jsonify({"ok": False, "message": "इस Google account को Admin access नहीं दिया गया है।"}), 403

    return jsonify({
        "ok": True,
        "email": email,
        "session_token": make_admin_session(email),
        "expires_in": ADMIN_SESSION_TTL,
        "message": "Admin login successful.",
    })

@app.get("/api/admin/session")
def admin_session():
    email, error = require_admin()
    if error:
        return error
    return jsonify({"ok": True, "email": email})

def dashboard_user_counts(rows, now=None):
    """Count people once across devices and older Telegram/mobile records."""
    now = (now or now_ist()).astimezone(IST)
    start = now.replace(hour=0, minute=0, second=0, microsecond=0)
    end = start + timedelta(days=1)
    people = {}
    def date(value):
        try:
            parsed = datetime.fromisoformat(str(value or '').replace('Z', '+00:00'))
            return (parsed.replace(tzinfo=IST) if parsed.tzinfo is None else parsed).astimezone(IST)
        except (ValueError, TypeError):
            return None
    for row in rows:
        key = normalise_mobile(row.get('mobile')) or clean(row.get('user_id')) or clean(row.get('visitor_id'))
        if not key:
            continue
        person = people.setdefault(key, {'signup': None, 'active_today': False})
        signup = date(row.get('signup_at'))
        if signup and (person['signup'] is None or signup < person['signup']):
            person['signup'] = signup
        # Signup opens the dashboard too. Refreshes do not add another person.
        visited = date(row.get('last_visit_at') or row.get('last_login_at'))
        if any(value and start <= value < end for value in (signup, visited)):
            person['active_today'] = True
    return {
        'total_users': len(people),
        'today_signups': sum(bool(p['signup'] and start <= p['signup'] < end) for p in people.values()),
        'today_active_users': sum(p['active_today'] for p in people.values()),
        'today_returning_users': sum(bool(p['active_today'] and p['signup'] and p['signup'] < start) for p in people.values()),
        'date_ist': now.date().isoformat(),
        'updated_at': now.isoformat(),
    }

_admin_stats_pool = ThreadPoolExecutor(max_workers=1, thread_name_prefix='admin-stats')
_admin_stats_lock = threading.Lock()
_admin_stats_future = None
_admin_stats_cache = None
_admin_stats_cache_at = 0
_admin_stats_scheduled = {}

def load_dashboard_user_stats():
    """Read storage on a background thread, never the API request worker."""
    remote = sheet_store.enabled()
    rows = sheet_store.call('list_visitors')['visitors'] if remote else []
    conn = user_db() if remote else visitor_db()
    try:
        rows = list(rows)
        if not remote:
            rows.extend(dict(row) for row in conn.execute('SELECT * FROM visitor_registrations').fetchall())
        for row in conn.execute('SELECT telegram_id,mobile,signup_at,last_login_at FROM users').fetchall():
            rows.append({**dict(row), 'visitor_id': 'telegram:' + str(row['telegram_id'])})
    finally:
        conn.close()
    return dashboard_user_counts(rows)

@app.get('/api/admin/stats')
def admin_dashboard_stats():
    global _admin_stats_future, _admin_stats_cache, _admin_stats_cache_at
    email, error = require_admin()
    if error:
        return error
    with _admin_stats_lock:
        if _admin_stats_future is not None and _admin_stats_future.done():
            future = _admin_stats_future
            _admin_stats_future = None
            _admin_stats_cache = future.result()
            _admin_stats_cache_at = time.monotonic()
        if _admin_stats_future is None and _admin_stats_cache is not None and request.args.get('refresh') != '1' and time.monotonic() - _admin_stats_cache_at < 60:
            response = jsonify({'ok': True, 'stats': _admin_stats_cache})
        else:
            if _admin_stats_future is None:
                _admin_stats_future = _admin_stats_pool.submit(load_dashboard_user_stats)
            response = jsonify({'ok': True, 'pending': True})
            response.status_code = 202
    response.headers['Cache-Control'] = 'no-store'
    return response

@app.post('/api/internal/admin-stats-refresh')
def scheduled_admin_stats_refresh():
    """A scoped signed job can refresh counts, never access admin/user data."""
    global _admin_stats_future, _admin_stats_cache, _admin_stats_cache_at
    timestamp = request.headers.get('X-Stats-Timestamp', '')
    signature = request.headers.get('X-Stats-Signature', '')
    secret = clean(os.getenv('TELEGRAM_BOT_TOKEN'))
    raw = request.get_data()
    if not secret or not timestamp.isdigit() or abs(time.time() - int(timestamp)) > 300 or len(raw) > 256:
        return jsonify({'ok': False, 'message': 'Scheduled job authentication required.'}), 401
    expected = hmac.new(secret.encode(), b'mp-admin-stats-refresh\n' + timestamp.encode() + b'\n' + raw, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(signature, expected):
        return jsonify({'ok': False, 'message': 'Scheduled job authentication required.'}), 401
    payload = request.get_json(silent=True) or {}
    today = now_ist()
    day, slot = payload.get('date_ist'), payload.get('slot')
    if day != today.date().isoformat() or slot not in ('07:00', '17:30') or slot > today.strftime('%H:%M'):
        return jsonify({'ok': False, 'message': 'Invalid or future update slot.'}), 400
    key = day + '@' + slot
    with _admin_stats_lock:
        for old in list(_admin_stats_scheduled):
            if not old.startswith(day + '@'):
                del _admin_stats_scheduled[old]
        future = _admin_stats_scheduled.get(key)
        if future is None:
            if _admin_stats_future is None or _admin_stats_future.done():
                _admin_stats_future = _admin_stats_pool.submit(load_dashboard_user_stats)
            future = _admin_stats_future
            _admin_stats_scheduled[key] = future
        if not future.done():
            response = jsonify({'ok': True, 'pending': True})
            response.status_code = 202
        else:
            try:
                stats = future.result()
            except Exception:
                del _admin_stats_scheduled[key]
                if _admin_stats_future is future:
                    _admin_stats_future = None
                raise
            if _admin_stats_cache is None or stats['updated_at'] >= _admin_stats_cache['updated_at']:
                _admin_stats_cache = stats
                _admin_stats_cache_at = time.monotonic()
            if _admin_stats_future is future:
                _admin_stats_future = None
            response = jsonify({'ok': True, 'completed': True, 'date_ist': day, 'slot': slot, 'updated_at': stats['updated_at']})
    response.headers['Cache-Control'] = 'no-store'
    return response


@app.post('/api/internal/signup-alert-status')
def signup_alert_status():
    """Signed setup check: report private bot readiness without user details."""
    import signup_alerts
    timestamp = request.headers.get('X-Alert-Timestamp', '')
    signature = request.headers.get('X-Alert-Signature', '')
    secret = clean(os.getenv('TELEGRAM_BOT_TOKEN'))
    raw = request.get_data()
    if not secret or not timestamp.isdigit() or abs(time.time() - int(timestamp)) > 300 or len(raw) > 256:
        return jsonify({'ok': False}), 401
    expected = hmac.new(secret.encode(), b'mp-signup-alert-status\n' + timestamp.encode() + b'\n' + raw, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(signature, expected):
        return jsonify({'ok': False}), 401
    if not signup_alerts.enabled():
        return jsonify({'ok': False, 'status': 'disabled'}), 503
    try:
        signup_alerts.private_destination(signup_alerts.configured_call, clean(os.getenv('TELEGRAM_ADMIN_CHAT_ID')))
    except Exception as exc:
        reason = 'bot_or_private_chat_not_ready' if isinstance(exc, ValueError) else type(exc).__name__
        return jsonify({'ok': False, 'status': reason}), 503
    return jsonify({'ok': True, 'bot_username': 'mptenders_bot', 'destination': 'private_admin',
                    'private_chat_verified': True, 'welcome_link': 'whatsapp_prefilled', 'outbox': 'saved_account_record'})

@app.post('/api/internal/welcome-backfill')
def welcome_backfill():
    import sys
    import welcome_backfill as batch
    import signup_alerts
    timestamp = request.headers.get('X-Alert-Timestamp', '')
    signature = request.headers.get('X-Alert-Signature', '')
    secret = clean(os.getenv('TELEGRAM_BOT_TOKEN'))
    raw = request.get_data()
    if not secret or not timestamp.isdigit() or abs(time.time()-int(timestamp)) > 300 or len(raw) > 512:
        return jsonify({'ok': False}), 401
    expected = hmac.new(secret.encode(), b'mp-welcome-backfill\n'+timestamp.encode()+b'\n'+raw, hashlib.sha256).hexdigest()
    if not hmac.compare_digest(signature, expected):
        return jsonify({'ok': False}), 401
    payload = request.get_json(silent=True) or {}
    if payload.get('campaign') != batch.CAMPAIGN or payload.get('cutoff') != batch.CUTOFF:
        return jsonify({'ok': False}), 400
    if not signup_alerts.enabled():
        return jsonify({'ok': False, 'status': 'disabled'}), 503
    try:
        signup_alerts.private_destination(signup_alerts.configured_call, clean(os.getenv('TELEGRAM_ADMIN_CHAT_ID')), use_cache=True)
    except Exception:
        return jsonify({'ok': False, 'status': 'private_chat_not_ready'}), 503
    status = batch.start(sys.modules[__name__], retry=payload.get('retry') is True)
    return jsonify({'ok': True, **status}), 200 if status['state'] == 'complete' else 202


@app.get("/api/admin/users-migration")
def users_migration_snapshot():
    """Private, consistent export for moving users and their login history."""
    email, error = require_admin()
    if error:
        return error
    conn = user_db()
    try:
        if conn.postgres:
            conn.execute("SET TRANSACTION ISOLATION LEVEL REPEATABLE READ, READ ONLY")
        else:
            conn.execute("BEGIN")
        users = [dict(row) for row in conn.execute("SELECT * FROM users ORDER BY telegram_id").fetchall()]
        events = [dict(row) for row in conn.execute("SELECT * FROM login_events ORDER BY id").fetchall()]
        return Response(json.dumps({"version": 1, "users": users, "login_events": events}),
                        mimetype="application/json", headers={"Cache-Control": "no-store",
                        "Content-Disposition": 'attachment; filename="mp-users-migration.private.json"'})
    finally:
        conn.close()

@app.post("/api/admin/logout")
def admin_logout():
    # Sessions are stateless and short-lived. Clearing the browser token is
    # sufficient; this endpoint exists for a clean client-side logout flow.
    return jsonify({"ok": True})

def github_dispatch(workflow, inputs=None):
    token = clean(os.getenv("GITHUB_ACTIONS_TOKEN"))
    if not token:
        raise RuntimeError("GITHUB_ACTIONS_TOKEN is not configured on Render.")
    owner = clean(os.getenv("GITHUB_REPO_OWNER", "esign2015"))
    repo = clean(os.getenv("GITHUB_REPO_NAME", "mp-tenders"))
    body = {"ref": clean(os.getenv("GITHUB_REPO_BRANCH", "main"))}
    if inputs:
        body["inputs"] = inputs
    url = f"https://api.github.com/repos/{owner}/{repo}/actions/workflows/{urllib_parse.quote(workflow, safe='')}/dispatches"
    data = json.dumps(body).encode("utf-8")
    req = urllib_request.Request(
        url,
        data=data,
        method="POST",
        headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/vnd.github+json",
            "X-GitHub-Api-Version": "2022-11-28",
            "Content-Type": "application/json",
            "User-Agent": "mp-tenders-admin",
        },
    )
    with urllib_request.urlopen(req, timeout=30) as response:
        if response.status not in (200, 201, 202, 204):
            raise RuntimeError(f"GitHub dispatch returned HTTP {response.status}")

def github_latest_workflow_run(workflow):
    token = clean(os.getenv("GITHUB_ACTIONS_TOKEN"))
    if not token:
        raise RuntimeError("GITHUB_ACTIONS_TOKEN is not configured on Render.")
    owner = clean(os.getenv("GITHUB_REPO_OWNER", "esign2015"))
    repo = clean(os.getenv("GITHUB_REPO_NAME", "mp-tenders"))
    url = (
        f"https://api.github.com/repos/{owner}/{repo}/actions/workflows/"
        f"{urllib_parse.quote(workflow, safe='')}/runs?event=workflow_dispatch&per_page=1"
    )
    req = urllib_request.Request(
        url,
        method="GET",
        headers={
            "Authorization": f"Bearer {token}",
            "Accept": "application/vnd.github+json",
            "X-GitHub-Api-Version": "2022-11-28",
            "User-Agent": "mp-tenders-admin",
        },
    )
    with urllib_request.urlopen(req, timeout=30) as response:
        payload = json.loads(response.read().decode("utf-8"))
    runs = payload.get("workflow_runs") or []
    if not runs:
        return {}
    run = runs[0]
    return {
        "id": run.get("id"),
        "run_number": run.get("run_number"),
        "status": run.get("status"),
        "conclusion": run.get("conclusion"),
        "created_at": run.get("created_at"),
        "updated_at": run.get("updated_at"),
        "html_url": run.get("html_url"),
        "head_sha": run.get("head_sha"),
    }

@app.get("/api/admin/workflow-status")
def admin_workflow_status():
    email, error = require_admin()
    if error:
        return error
    action = clean(request.args.get("action")).lower()
    mapping = {
        "data_refresh": "scrape.yml",
        "retry_pending": "targeted-pending-retry.yml",
        "telegram_test": "telegram-test.yml",
        "closing_today": "telegram_manual_pdf.yml",
        "closing_tomorrow": "telegram_manual_pdf.yml",
        "new_today": "telegram_manual_pdf.yml",
    }
    workflow = mapping.get(action)
    if not workflow:
        return jsonify({"ok": False, "message": "Unknown workflow status request."}), 400
    try:
        return jsonify({"ok": True, "action": action, "workflow": workflow, "run": github_latest_workflow_run(workflow), "requested_by": email})
    except Exception as exc:
        return jsonify({"ok": False, "message": str(exc)}), 502


@app.post("/api/admin/action")
def admin_action():
    email, error = require_admin()
    if error:
        return error
    payload = request.get_json(silent=True) or {}
    action = clean(payload.get("action")).lower()
    try:
        if action == "telegram_pdf":
            report = clean(payload.get("report", "closing_today"))
            if report not in {"closing_today", "closing_tomorrow", "new_today"}:
                return jsonify({"ok": False, "message": "Invalid PDF report."}), 400
            view = clean(payload.get("view", "table")).lower()
            if view != "table":
                return jsonify({"ok": False, "message": "Invalid PDF view."}), 400
            github_dispatch("telegram_manual_pdf.yml", {"report": report, "view": view})
            message = f"Telegram {view.title()} PDF workflow started: {report}"
        elif action == "data_refresh":
            github_dispatch("scrape.yml")
            message = "Full data refresh workflow started."
        elif action == "retry_pending":
            github_dispatch("targeted-pending-retry.yml")
            message = "Pending detail retry workflow started."
        elif action == "telegram_test":
            github_dispatch("telegram-test.yml")
            message = "Telegram test workflow started."
        else:
            return jsonify({"ok": False, "message": "Unknown admin command."}), 400
        from bidder_tools import queue_system_audit
        queue_system_audit(sys.modules[__name__],email,action,report=payload.get("report",""))
        return jsonify({"ok": True, "message": message, "requested_by": email})
    except Exception as exc:
        print(f"Admin action failed: {exc}")
        return jsonify({"ok": False, "message": str(exc)}), 502

@app.get("/health")
def health():
    return jsonify({
        "status": "healthy",
        "csv_exists": CSV_FILE.exists(),
        "records": len(read_rows()),
        "user_db_backend": user_db_backend(),
        "updated_at": datetime.now(timezone.utc).isoformat(),
    })


@app.get("/api/tenders")
def tenders():
    rows=read_rows()
    try:
        from bidder_tools import public_json
        metadata=public_json("inventory_counts.json")
    except Exception:metadata={}
    return jsonify({
        "updated_at": metadata.get("snapshot_at"),
        "count": len(rows),
        "tenders": rows,
    })


@app.post("/api/fetch")
def fetch():
    return jsonify(scrape_mp_tenders(CSV_FILE))


import sys
from account_access import install as install_account_access
account_service=install_account_access(sys.modules[__name__])
from bidder_tools import install as install_bidder_tools
install_bidder_tools(sys.modules[__name__])

if __name__ == "__main__":
    app.run(host="0.0.0.0", port=int(os.getenv("PORT", "10000")))
