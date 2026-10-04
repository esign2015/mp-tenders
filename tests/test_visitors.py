import sys,tempfile,uuid,unittest
from pathlib import Path
from unittest.mock import patch
from io import BytesIO
from openpyxl import load_workbook
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'backend'))
import server

class VisitorTests(unittest.TestCase):
    def setUp(self):
        self.tmp=tempfile.TemporaryDirectory()
        self.database=patch.object(server,'DATABASE_URL','');self.database.start()
        self.path=patch.object(server,'USER_DB_PATH',Path(self.tmp.name)/'users.db');self.path.start()
        self.secret=patch.dict('os.environ',{'VISITOR_SESSION_SECRET':'visitor-test-secret','VISITOR_PROFILE_LEGACY_ALLOWED':'1'});self.secret.start()
        self.client=server.app.test_client()
    def tearDown(self):
        self.secret.stop();self.path.stop();self.database.stop();self.tmp.cleanup()
    def payload(self):return {'registration_id':str(uuid.uuid4()),'name':'Visitor Test','mobile':'9876543210','district':'Dewas'}
    def test_save_three_fields_once_restore_and_export_without_telegram(self):
        data=self.payload()
        with patch.object(server,'telegram_api',side_effect=AssertionError('Telegram must not be called')):
            response=self.client.post('/api/visitors/register',json=data)
            self.assertEqual(response.status_code,200)
            self.assertEqual(response.json['profile'],{'name':'Visitor Test','mobile':'+919876543210','district':'Dewas'})
            self.assertEqual(self.client.post('/api/visitors/register',json=data).status_code,200)
            token=response.json['session_token']
            restored=self.client.post('/api/visitors/session',json={'session_token':token})
            self.assertEqual(restored.status_code,200)
            self.assertEqual(restored.headers['Cache-Control'],'no-store')
        self.assertIsNone(server.read_telegram_session(token))
        conn=server.visitor_db()
        self.assertEqual(conn.execute('SELECT COUNT(*) FROM visitor_registrations').fetchone()[0],1)
        self.assertEqual(conn.execute('SELECT COUNT(*) FROM users').fetchone()[0],0)
        self.assertEqual(conn.execute('SELECT visit_count FROM visitor_registrations').fetchone()[0],2)
        conn.close()
        workbook=load_workbook(BytesIO(server.build_user_excel()))
        self.assertEqual(list(workbook['Visitor Registrations'].values)[1][1:4],('Visitor Test','+919876543210','Dewas'))
    def test_invalid_input_and_tampered_token_are_rejected(self):
        for change in ({'name':''},{'mobile':'1234'},{'district':''},{'registration_id':'1'}):
            self.assertEqual(self.client.post('/api/visitors/register',json={**self.payload(),**change}).status_code,400)
        token=self.client.post('/api/visitors/register',json=self.payload()).json['session_token']
        self.assertEqual(self.client.post('/api/visitors/session',json={'session_token':token+'x'}).status_code,401)
        self.assertEqual(self.client.post('/api/users/export',json={'session_token':token}).status_code,403)
        self.assertEqual(self.client.get('/api/admin/visitor-registrations').status_code,401)
    def test_replayed_registration_cannot_overwrite_profile(self):
        data=self.payload();self.client.post('/api/visitors/register',json=data)
        self.assertEqual(self.client.post('/api/visitors/register',json={**data,'name':'Different person'}).status_code,409)
    def test_formula_input_is_exported_as_text(self):
        self.client.post('/api/visitors/register',json={**self.payload(),'name':'=1+1'})
        workbook=load_workbook(BytesIO(server.build_user_excel()))
        self.assertEqual(workbook['Visitor Registrations']['B2'].value,'=1+1')
        self.assertEqual(workbook['Visitor Registrations']['B2'].data_type,'s')
    def test_document_profile_requires_email_and_validates_registration_dates(self):
        token=self.client.post('/api/visitors/register',json=self.payload()).json['session_token']
        profile={'bidderName':'Bidder','firmName':'Firm','status':'Proprietor','place':'Dewas','relative':'no','email':'bidder@example.test'}
        for change in ({'email':''},{'email':'invalid'},{'representativeEmail':'bad address'},{'registrationDate':'2026-02-30'},{'pincode':'12345'},{'pincode':'000000'}):
            response=self.client.post('/api/visitors/affidavit',json={'session_token':token,'profile':{**profile,**change}})
            self.assertEqual(response.status_code,400)
        profile.update(pincode='455332',district='Dewas',tehsil='Kannod',pan='ABCDE1234F',gst='23ABCDE1234F1Z5',registrationNumber='REG-42',registrationDate='2025-01-01',registrationValidTill='2030-01-01',representativeName='Different Person',representativeEmail='rep@example.test')
        self.assertEqual(self.client.post('/api/visitors/affidavit',json={'session_token':token,'profile':profile}).status_code,200)
        restored=self.client.post('/api/visitors/affidavit',json={'session_token':token}).json['profile']
        for key in profile:self.assertEqual(restored[key],profile[key])

    def test_structured_address_and_gst_pan_validation(self):
        token=self.client.post('/api/visitors/register',json=self.payload()).json['session_token']
        profile={'bidderName':'Bidder','firmName':'Firm','status':'Proprietor','place':'Kannod','relative':'no','email':'bidder@example.test','houseNumber':'12','roadStreet':'Main Road','locality':'Pipla','landmark':'Near Temple','gst':'23abcde1234f1z5','pan':''}
        saved=self.client.post('/api/visitors/affidavit',json={'session_token':token,'profile':profile})
        self.assertEqual(saved.status_code,200)
        self.assertEqual(saved.json['profile']['address'],'12, Main Road, Pipla, Near Temple')
        self.assertEqual(saved.json['profile']['gst'],'23ABCDE1234F1Z5')
        self.assertEqual(saved.json['profile']['pan'],'ABCDE1234F')
        restored=self.client.post('/api/visitors/affidavit',json={'session_token':token}).json['profile']
        for key in ('houseNumber','roadStreet','locality','landmark'):self.assertEqual(restored[key],profile[key])
        for change in ({'gst':'23ABCDE1234F1Z'},{'gst':'AAABCDE1234F1Z5'},{'gst':'231BCDE1234F1Z5'},{'pan':'ABCDE12345'},{'pan':'ABCD11234F'},{'houseNumber':'','roadStreet':'','locality':'','landmark':''}):
            self.assertEqual(self.client.post('/api/visitors/affidavit',json={'session_token':token,'profile':{**profile,**change}}).status_code,400)
        # Valid exceptions can keep a different PAN; old cached clients may still edit the full address.
        self.assertEqual(self.client.post('/api/visitors/affidavit',json={'session_token':token,'profile':{**profile,'pan':'PQRST6789Z'}}).status_code,200)
        old_client={k:v for k,v in profile.items() if k not in ('houseNumber','roadStreet','locality','landmark')}
        changed=self.client.post('/api/visitors/affidavit',json={'session_token':token,'profile':{**old_client,'address':'Updated legacy address'}})
        self.assertEqual(changed.status_code,200);self.assertEqual(changed.json['profile']['roadStreet'],'Updated legacy address')
        self.assertEqual(changed.json['profile']['houseNumber'],'')

    def test_affidavit_profile_restores_only_in_its_signed_session(self):
        first=self.client.post('/api/visitors/register',json=self.payload()).json
        second=self.client.post('/api/visitors/register',json=self.payload()).json
        profile={'bidderName':'Bidder','email':'bidder@example.test','firmName':'Firm','status':'Proprietor','place':'Dewas','relative':'no'}
        saved=self.client.post('/api/visitors/affidavit',json={'session_token':first['session_token'],'profile':profile})
        self.assertEqual(saved.status_code,200)
        restored=self.client.post('/api/visitors/session',json={'session_token':first['session_token']})
        self.assertEqual(restored.json['affidavit_profile']['firmName'],'Firm')
        isolated=self.client.post('/api/visitors/affidavit',json={'session_token':second['session_token']})
        self.assertEqual(isolated.json['profile'],{})
        self.assertEqual(self.client.post('/api/visitors/affidavit',json={'session_token':first['session_token']+'x'}).status_code,401)

if __name__=='__main__':unittest.main()
